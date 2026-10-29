"""Genera el PowerPoint del proyecto desde las métricas y gráficos exportados por Colab.
Uso: python src/crear_presentacion.py
Para PDF: libreoffice -env:UserInstallation=file:///tmp/lo_proj --headless --convert-to pdf --outdir presentacion presentacion/*.pptx
"""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT=Path(__file__).resolve().parents[1]
M=json.loads((ROOT/'resultados/metricas_modelo.json').read_text(encoding='utf-8'))
G=ROOT/'resultados/graficos'
P=ROOT/'presentacion'/'Proyecto_Final_Analisis_Salarios.pptx'
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
NAVY=RGBColor(12,26,45); NAVY2=RGBColor(20,40,62)
TEAL=RGBColor(28,190,160); WHITE=RGBColor(246,250,253); MUTED=RGBColor(174,195,211)
BLUE=RGBColor(60,133,231); PURPLE=RGBColor(156,111,234); ORANGE=RGBColor(245,177,78)
DARK_TEXT=RGBColor(35,53,70)

def rect(s,x,y,w,h,color,radius=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
 sh.fill.solid();sh.fill.fore_color.rgb=color; sh.line.fill.background()
 return sh

def text(s,tx,x,y,w,h,size=22,color=WHITE,bold=False,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 tf=sh.text_frame;tf.word_wrap=True
 tf.margin_left=Inches(.02);tf.margin_right=Inches(.02);tf.margin_top=Inches(.01);tf.margin_bottom=Inches(.01)
 for j,line in enumerate(tx.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
  p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=color
  p.space_after=Pt(8)
  if align is not None:p.alignment=align
 return sh

def slide(title,kicker='CIENCIA DE DATOS EN ACCIÓN',number=True):
 s=prs.slides.add_slide(prs.slide_layouts[6]);rect(s,0,0,13.333,7.5,NAVY)
 rect(s,0,0,.10,7.5,TEAL);text(s,kicker.upper(),.52,.27,12,.25,11,TEAL,True)
 text(s,title,.52,.72,12.3,.8,29,WHITE,True)
 if number:
  rect(s,.52,7.04,12.25,.012,RGBColor(53,75,94))
  text(s,f'PROYECTO FINAL · DATOS SINTÉTICOS                                      {len(prs.slides):02d} / 11',.52,7.15,12.3,.20,9,MUTED)
 return s

def panel(s,x,y,w,h,heading,body,accent=TEAL,fs=16):
 rect(s,x,y,w,h,NAVY2,True); rect(s,x+.18,y+.22,.06,.38,accent)
 text(s,heading,x+.40,y+.22,w-.60,.43,19,WHITE,True)
 text(s,body,x+.32,y+.80,w-.64,h-.90,fs,MUTED)

def picture(s,name,x,y,w,h):
 from PIL import Image
 img=Image.open(G/name);iw,ih=img.size; r=iw/ih;boxr=w/h
 if r>boxr: realw=w;realh=w/r;xx=x;yy=y+(h-realh)/2
 else: realh=h;realw=h*r;xx=x+(w-realw)/2;yy=y
 s.shapes.add_picture(str(G/name), Inches(xx), Inches(yy), width=Inches(realw),height=Inches(realh))

s=slide('Análisis y predicción de salarios',kicker='PROYECTO FINAL · INGENIERÍA DE SISTEMAS')
text(s,'De los datos a decisiones responsables',.62,1.85,11,.65,28,MUTED)
rect(s,.62,2.77,10.7,.02,TEAL)
text(s,'Exploración  ·  Estadística  ·  Regresión lineal  ·  Storytelling',.62,3.23,11,.7,22,WHITE)
text(s,'Equipo: [Integrante 1] · [Integrante 2] · [Integrante 3 opcional]',.62,5.65,11,.40,17,MUTED)
text(s,'Dataset ficticio de nómina en Colombia · Octubre 2026',.62,6.20,11,.32,15,TEAL)

s=slide('El problema y la pregunta de investigación')
panel(s,.58,1.73,6.05,4.55,'CONTEXTO PROFESIONAL','Un área de Talento Humano necesita entender cómo varían los salarios y evaluar si un modelo puede estimar valores orientativos a partir de información laboral. El análisis no reemplaza la política salarial.',TEAL,18)
panel(s,6.85,1.73,5.9,4.55,'PREGUNTA CENTRAL','¿Cómo se relacionan el nivel de cargo, la experiencia y el departamento con el salario mensual? ¿Una regresión lineal supera una predicción basada solo en la mediana?',BLUE,18)

s=slide('Etapa 1 · Origen y estructura de los datos')
for x,num,lab,acc in [(0.6,'1.514','filas originales',TEAL),(4.77,'12','columnas',BLUE),(8.93,'100%','datos ficticios',PURPLE)]:
 rect(s,x,1.86,3.8,1.75,NAVY2,True);text(s,num,x+.27,2.04,3.2,.8,34,acc,True);text(s,lab,x+.27,2.92,3.3,.35,16,MUTED)
text(s,'Variables: salario mensual COP, nivel de cargo, experiencia, antigüedad, departamento, educación, ciudad, modalidad y horas extra.',.72,4.22,11.6,1.10,21,WHITE)
text(s,'Fuente: generación reproducible con Python (semilla 202611). No contiene registros reales, datos personales ni información empresarial.',.72,5.63,11.7,.75,16,MUTED)

s=slide('Etapa 1 · Depuración y control de calidad')
for i,(n,label,accent) in enumerate([('1.514','Registros originales',BLUE),('−14','Duplicados exactos',ORANGE),('−16','Salarios/edades inválidos',ORANGE),('1.484','Registros finales',TEAL)]):
 x=.72+i*3.18;rect(s,x,2.03,2.88,2.26,NAVY2,True);text(s,n,x+.22,2.47,2.35,.65,30,accent,True);text(s,label,x+.18,3.25,2.57,.72,15,MUTED)
text(s,'Además: se estandarizaron etiquetas, se imputaron valores faltantes y se conservaron outliers plausibles.',.78,5.15,11.8,1.1,20,WHITE)

s=slide('Etapa 2 · ¿Cómo se distribuyen los salarios?')
picture(s,'01_distribucion_salarios.png',.55,1.65,8.24,4.93)
panel(s,9.07,1.72,3.57,4.80,'LECTURA',f'Media: ${M["salario_promedio_cop"]:,.0f}\n\nMediana: ${M["salario_mediano_cop"]:,.0f}\n\nSe observa dispersión entre salarios; no todos ganan cerca del promedio.',TEAL,16)

s=slide('Etapa 2 · Comparación por grupo')
picture(s,'02_salario_por_departamento.png',.54,1.82,6.08,4.55)
picture(s,'03_boxplot_nivel_cargo.png',6.80,1.82,6.00,4.55)
text(s,'Los promedios departamentales también reflejan la composición de cargos.',.65,6.43,12,.37,17,MUTED)

s=slide('Etapa 2 · Relaciones e hipótesis')
picture(s,'04_matriz_correlaciones.png',.55,1.72,7.25,4.91)
panel(s,8.05,1.82,4.62,4.68,'HIPÓTESIS',f'Nivel de cargo ↔ salario: r = {M["correlacion_nivel_salario"]:.2f}.\n\nExperiencia ↔ salario: r = {M["correlacion_experiencia_salario"]:.2f}.\n\nCorrelación no significa causalidad.',PURPLE,17)

s=slide('Etapa 3 · Diseño de la predicción')
panel(s,.58,1.78,3.68,4.77,'1 · PREPARAR','Predictores: cargo, experiencia, antigüedad, departamento, educación, ciudad, modalidad y horas extra. Transformación categórica One Hot.',TEAL,17)
panel(s,4.81,1.78,3.70,4.77,'2 · ENTRENAR',f'{M["n_entrenamiento"]:,} registros (80%) para ajustar regresión lineal. Se compara con el valor mediano del entrenamiento.',BLUE,17)
panel(s,9.04,1.78,3.70,4.77,'3 · EVALUAR',f'{M["n_prueba"]:,} registros (20%) reservados. Medimos MAE, RMSE, R² y validación cruzada de 5 pliegues.',PURPLE,17)

s=slide('Etapa 3 · ¿Qué tan bien funcionó?')
picture(s,'06_real_vs_predicho.png',.53,1.62,7.52,4.88)
panel(s,8.32,1.69,4.34,4.85,'RESULTADOS',f'R² en prueba: {M["regresion"]["R2"]:.3f}\n\nMAE del modelo: ${M["regresion"]["MAE_COP"]:,.0f}\n\nMAE del modelo base: ${M["baseline"]["MAE_COP"]:,.0f}\n\nEl modelo mejora al baseline.',TEAL,16)

s=slide('Storytelling · De un CSV a una decisión informada')
panel(s,.58,1.78,5.88,4.77,'LO QUE APRENDIMOS','La limpieza evita conclusiones engañosas. La estadística resume la dispersión. El modelo necesita compararse con una línea base y evaluarse sobre datos nunca vistos.',BLUE,18)
panel(s,6.72,1.78,6.01,4.77,'APLICACIÓN PROPUESTA','Un prototipo que valide extractos autorizados de SAP HCM y muestre patrones agregados para planeación presupuestaria. Sin decisiones individuales automáticas.',TEAL,18)

s=slide('Conclusiones, límites y próximos pasos')
panel(s,.56,1.75,4.02,4.90,'CONCLUSIÓN','Se construyó un flujo completo, reproducible y evaluado. El modelo explica parte importante de la variación presente en la simulación.',TEAL,18)
panel(s,4.70,1.75,4.02,4.90,'LÍMITES','Datos sintéticos; sin generalización automática a salarios reales. No inferimos causalidad. Un outlier no es siempre un dato erróneo.',ORANGE,18)
panel(s,8.84,1.75,3.92,4.90,'PRÓXIMO PASO','Con permisos y datos anonimizados: validar calidad, vigilar sesgos y contrastar predicciones con especialistas humanos.',PURPLE,18)

P.parent.mkdir(exist_ok=True)
prs.save(P)
print('Presentación lista:',P, 'diapositivas',len(prs.slides))
